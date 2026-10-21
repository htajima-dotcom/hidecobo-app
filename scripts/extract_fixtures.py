"""既存Excel（社労士準拠版）から給与計算の入力と期待値を取り出し、テスト用JSONにする。

使い方: python3 -I scripts/extract_fixtures.py <xlsxファイル...> <出力ディレクトリ>
計算済みの値（キャッシュ）があるシートだけを出力する。
"""
import json
import os
import sys
from datetime import date, datetime

import openpyxl

# 給与計算_X月度 の列（0始まり）
PAY_COLS = {
    "no": 0, "name": 1, "contract": 2, "payType": 3, "hourlyRate": 4,
    "workDays": 5, "scheduledHours": 6, "paidLeaveDays": 7, "holidays": 8,
    "overtimeHours": 9, "commuteDays": 10,
    "executivePay": 11, "basePay": 12, "paidLeavePay": 13, "positionAllowance": 14,
    "adjustmentAllowance": 15, "overtimeRate": 16, "overtimePay": 17,
    "otherTaxable": 18, "commuteAdjustment": 19, "commutePay": 20,
    "commuteNonTaxable": 21, "commuteTaxable": 22, "taxableTotal": 23,
    "nonTaxableTotal": 24, "grossPay": 25, "healthInsurance": 26,
    "careInsurance": 27, "childSupport": 28, "pension": 29,
    "employmentInsurance": 30, "socialInsuranceTotal": 31, "afterSocialInsurance": 32,
    "dependents": 33, "taxColumn": 34, "incomeTaxTable": 35, "incomeTaxManual": 36,
    "incomeTax": 37, "residentTax": 38, "housing": 39, "otherDeduction": 40,
    "deductionTotal": 41, "netPay": 42, "salarySubtotal": 43, "attendanceName": 44,
    "note": 45,
}
# 日曜出勤の3列はファイルによって位置が違うので、見出しで探す
SUNDAY_HEADERS = {"sundayHours": "日曜出勤(時間)", "sundayRate": "休日単価(1.35倍)",
                  "sundayPay": "休日出勤手当(日曜)"}
INPUT_KEYS = ["workDays", "scheduledHours", "paidLeaveDays", "holidays", "overtimeHours",
              "commuteDays", "sundayHours", "otherTaxable", "commuteAdjustment",
              "incomeTaxManual", "otherDeduction"]
OUTPUT_KEYS = ["hourlyRate", "executivePay", "basePay", "paidLeavePay", "positionAllowance",
               "adjustmentAllowance", "overtimeRate", "overtimePay", "commutePay",
               "commuteNonTaxable", "commuteTaxable", "taxableTotal", "nonTaxableTotal",
               "grossPay", "healthInsurance", "careInsurance", "childSupport", "pension",
               "employmentInsurance", "socialInsuranceTotal", "afterSocialInsurance",
               "incomeTaxTable", "incomeTax", "residentTax", "housing", "deductionTotal",
               "netPay", "salarySubtotal", "sundayRate", "sundayPay"]

# 給与計算設定 の従業員行（A〜AA）
EMP_COLS = ["no", "name", "attendanceName", "contract", "payType", "executivePay",
            "monthlyBase", "hourlyRate", "newHourlyRate", "newHourlyRateFrom",
            "paidLeaveHoursPerDay", "positionAllowance", "adjustmentAllowance",
            "adjustmentInOvertime", "commuteMethod", "commuteRoundTripKm", "commuteFixed",
            "commuteTaxClass", "healthStandard", "pensionStandard", "careInsured",
            "employmentInsured", "taxColumn", "dependents", "residentTax", "residentTaxJune",
            "housing"]


def plain(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def num_or_none(v):
    return None if v in (None, "") else plain(v)


def read_rules(s):
    g = lambda c: plain(s[c].value)
    mileage = []
    for r in range(6, 18):
        km, cap = s[f"L{r}"].value, s[f"M{r}"].value
        if km is not None:
            mileage.append({"minOneWayKm": plain(km), "limit": plain(cap)})
    months = {}
    for r in range(6, 20):
        label = s[f"O{r}"].value
        if not label or not str(label).endswith("月度"):
            continue
        months[label] = {
            "payDate": g(f"P{r}"), "taxTable": g(f"Q{r}"),
            "healthRate": g(f"R{r}"), "careRate": g(f"S{r}"), "childSupportRate": g(f"T{r}"),
            "pensionRate": g(f"U{r}"), "employmentRate": g(f"V{r}"),
            "deductSocialInsurance": g(f"W{r}") == 1,
        }
    rules = {
        "monthlyScheduledHours": round((365 - g("E5")) * 8 / 12, 2),
        "overtimeMultiplier": g("E9"),
        "fuelEfficiencyKmPerL": g("E10"),
        "gasolinePricePerL": g("E11"),
        "commuteCap": g("E12"),
        "hourlyRounding": g("E16"),
        "sundayMultiplier": g("E19"),
        "mileageNonTaxableLimits": mileage,
    }
    return rules, months


def read_employees(s):
    emps = []
    for r in range(31, 80):
        vals = [s.cell(row=r, column=c + 1).value for c in range(len(EMP_COLS))]
        if vals[1] in (None, ""):
            continue
        e = {k: plain(v) for k, v in zip(EMP_COLS, vals)}
        e["no"] = str(e["no"]) if e["no"] is not None else None
        e["row"] = r
        # 標準報酬の改定（随時改定・定時決定）：AF〜AJ列
        rev = [plain(s.cell(row=r, column=c).value) for c in range(32, 37)]
        if (rev[0] or rev[1]) and rev[2]:
            e["standardRevision"] = {
                "healthStandard": rev[0], "pensionStandard": rev[1], "fromPayDate": rev[2],
                "settleOnPayDate": rev[3], "settleMonths": rev[4] or 0}
        emps.append(e)
    return emps


def read_tax_tables(s):
    def table(lo_col, first, last, hi_first, hi_last, otsu_col):
        cols = lambda start: [openpyxl.utils.get_column_letter(start + i) for i in range(8)]
        base = openpyxl.utils.column_index_from_string(lo_col)
        rows = []
        for r in range(first, last + 1):
            rows.append({
                "min": plain(s.cell(row=r, column=base).value),
                "max": plain(s.cell(row=r, column=base + 1).value),
                "kou": [plain(s.cell(row=r, column=base + 2 + i).value) for i in range(8)],
                "otsu": plain(s[f"{otsu_col}{r}"].value),
            })
        high = []
        for r in range(hi_first, hi_last + 1):
            high.append({
                "min": plain(s.cell(row=r, column=base).value),
                "rate": plain(s.cell(row=r, column=base + 1).value),
                "kou": [plain(s.cell(row=r, column=base + 2 + i).value) for i in range(8)],
            })
        return {"rows": rows, "high": high}

    r8 = table("A", 5, 236, 240, 248, "K")
    r9 = table("N", 5, 233, 237, 245, "X")
    # 乙欄の74万円以上は税額表の式（ExcelのAJ列の定数）をそのまま持つ
    r8.update({"taxFreeBelow": 105000, "otsuHigh": [
        {"min": 740000, "base": 259200, "rate": 0.4084},
        {"min": 1710000, "base": 655400, "rate": 0.45945}]})
    r9.update({"taxFreeBelow": 111000, "otsuHigh": [
        {"min": 740000, "base": 259000, "rate": 0.4084},
        {"min": 1720000, "base": 659300, "rate": 0.45945}]})
    return {"令和8年分": r8, "令和9年分": r9}


def main(paths, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    tax_written = False
    for path in paths:
        wbf = openpyxl.load_workbook(path, read_only=False)
        wbv = openpyxl.load_workbook(path, read_only=False, data_only=True)
        if not tax_written:
            with open(os.path.join(out_dir, "tax-tables.json"), "w") as f:
                json.dump(read_tax_tables(wbv["源泉徴収税額表"]), f, ensure_ascii=False, indent=1)
            tax_written = True
        rules, months = read_rules(wbv["給与計算設定"])
        emps = read_employees(wbv["給与計算設定"])
        for sheet in [n for n in wbv.sheetnames if n.startswith("給与計算_")]:
            label = sheet.split("_")[1]
            sv = wbv[sheet]
            header = [str(c.value or "").replace("\n", "") for c in sv[5]]
            cols = dict(PAY_COLS)
            for key, text in SUNDAY_HEADERS.items():
                cols[key] = header.index(text)
            rows = list(sv.iter_rows(min_row=6, max_row=6 + len(emps) - 1, values_only=True))
            if rows[0][PAY_COLS["netPay"]] is None:
                print(f"skip {os.path.basename(path)} {sheet}: 計算済みの値なし")
                continue
            sf = wbf[sheet]
            cases = []
            # 標準報酬の改定を反映する列がない月度（改定を知る前に確定した月）は改定を渡さない
            has_revision = "適用する健保標準報酬" in header
            for i, (e, row) in enumerate(zip(emps, rows)):
                if not has_revision:
                    e = {k: v for k, v in e.items() if k != "standardRevision"}
                get = lambda k: row[cols[k]] if cols[k] < len(row) else None
                assert get("name") == e["name"], (sheet, get("name"), e["name"])
                inp = {k: num_or_none(get(k)) for k in INPUT_KEYS}
                # 調整手当（遡及等）は式を上書きして手入力されていることがある
                adj = sf.cell(row=6 + i, column=PAY_COLS["adjustmentAllowance"] + 1).value
                if adj is not None and not str(adj).startswith("="):
                    inp["adjustmentOverride"] = plain(adj)
                cases.append({
                    "employee": e,
                    "input": inp,
                    "expected": {k: num_or_none(get(k)) for k in OUTPUT_KEYS},
                })
            payload = {"source": os.path.basename(path), "sheet": sheet, "month": label,
                       "rules": rules, "monthSettings": months[label], "cases": cases}
            pay_date = months[label]["payDate"]
            with open(os.path.join(out_dir, f"payroll-{pay_date[:7]}.json"), "w") as f:
                json.dump(payload, f, ensure_ascii=False, indent=1)
            print(f"wrote {pay_date[:7]} ({sheet}, {len(cases)}人)")


if __name__ == "__main__":
    main(sys.argv[1:-1], sys.argv[-1])
