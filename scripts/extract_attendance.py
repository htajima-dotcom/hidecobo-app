"""既存Excelの月度シート（日ごとの区分・出退勤時刻）を取り出し、勤怠集計のテスト用JSONにする。

期待値は「給与計算_X月度」の勤怠列（出勤日数・所定時間・有休・公休・残業・通勤日数・日曜出勤）。
使い方: python3 -I scripts/extract_attendance.py <xlsxファイル...> <出力ディレクトリ>
"""
import json
import os
import re
import sys
from datetime import date, datetime, time, timedelta

import openpyxl

EXPECTED_COLS = {"workDays": 5, "scheduledHours": 6, "paidLeaveDays": 7, "holidays": 8,
                 "overtimeHours": 9, "commuteDays": 10}


def hhmm(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        v = v.time()
    if isinstance(v, time):
        return f"{v.hour:02d}:{v.minute:02d}" + (f":{v.second:02d}" if v.second else "")
    if isinstance(v, timedelta):
        s = int(round(v.total_seconds()))
        return f"{s // 3600:02d}:{s % 3600 // 60:02d}"
    if isinstance(v, (int, float)):
        m = round(v * 1440)
        return f"{m // 60:02d}:{m % 60:02d}"
    return str(v)


def num(v):
    if v in (None, ""):
        return 0
    return int(v) if isinstance(v, float) and v.is_integer() else v


def main(paths, out_dir):
    for path in paths:
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        wbf = openpyxl.load_workbook(path, read_only=False)
        month_sheets = [n for n in wb.sheetnames if n.endswith(")") and "月度 (" in n]
        for ms in month_sheets:
            label = ms.split(" ")[0]
            pay = wb[f"給与計算_{label}"] if f"給与計算_{label}" in wb.sheetnames else None
            if pay is None or pay.cell(row=6, column=43).value is None:
                print(f"skip {ms}: 計算済みの値なし")
                continue
            ws = wb[ms]
            wf = wbf[ms]
            master = wb["マスタ・給与設定"]
            dates = []
            for c in range(4, 35):
                d = ws.cell(row=4, column=c).value
                if isinstance(d, datetime):
                    dates.append((c, d.date().isoformat()))
            header = [str(c.value or "").replace("\n", "") for c in pay[5]]
            sunday_col = header.index("日曜出勤(時間)")
            expected = {}
            for row in pay.iter_rows(min_row=6, max_row=40, values_only=True):
                name = row[44] if len(row) > 44 else None
                if name:
                    exp = {k: num(row[i]) for k, i in EXPECTED_COLS.items()}
                    exp["sundayHours"] = num(row[sunday_col])
                    expected[name] = exp
            people = []
            r = 6
            while r <= ws.max_row:
                name = ws.cell(row=r, column=2).value
                if ws.cell(row=r, column=3).value == "区分" and name:
                    rules = {}
                    # 休憩の自動判定（12時前出勤・13時後退勤で1時間）を使わない人
                    rest_formula = str(wf.cell(row=r + 3, column=4).value or "")
                    if rest_formula.startswith("=") and "TIME(12,0,0)" not in rest_formula:
                        rules["noBreak"] = True
                    # 計算開始時刻（これより前の出勤は数えない）
                    m = re.search(r"マスタ・給与設定!\$S\$(\d+)", str(wf.cell(row=r + 4, column=4).value or ""))
                    if m:
                        start = hhmm(master.cell(row=int(m.group(1)), column=19).value)
                        if start and start != "00:00":
                            rules["calcStart"] = start
                    days = []
                    for c, d in dates:
                        status = ws.cell(row=r, column=c).value
                        day = {"date": d, "status": status or "",
                               "in": hhmm(ws.cell(row=r + 1, column=c).value),
                               "out": hhmm(ws.cell(row=r + 2, column=c).value)}
                        rest = wf.cell(row=r + 3, column=c).value
                        if rest is not None and not str(rest).startswith("="):
                            h, mi = hhmm(rest).split(":")[:2]
                            day["breakMinutes"] = int(h) * 60 + int(mi)
                        days.append(day)
                    if name in expected:
                        people.append({"name": name, "rules": rules, "days": days,
                                       "expected": expected[name]})
                    r += 6
                else:
                    r += 1
            pay_month = (datetime.fromisoformat(dates[-1][1]).replace(day=1)).strftime("%Y-%m")
            with open(os.path.join(out_dir, f"attendance-{pay_month}.json"), "w") as f:
                json.dump({"source": os.path.basename(path), "sheet": ms, "people": people},
                          f, ensure_ascii=False, indent=1)
            print(f"wrote attendance-{pay_month} ({ms}, {len(people)}人)")


if __name__ == "__main__":
    main(sys.argv[1:-1], sys.argv[-1])
